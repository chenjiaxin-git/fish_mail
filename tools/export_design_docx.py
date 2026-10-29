"""Generate the Word design strictly from docs/开发设计文档.md. Requires python-docx."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
root=Path(__file__).resolve().parents[1]
source=root/'docs/开发设计文档.md';target=source.with_suffix('.docx')
def clean(s):
 s=re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1（\2）',s)
 return s.replace('**','').replace('`','').replace('\\|','|').strip()
def cells(line):
 result=[];current='';code=False;escape=False
 for ch in line.strip().strip('|'):
  if ch=='`' and not escape:code=not code
  if ch=='|' and not code and not escape:result.append(clean(current));current=''
  else:current+=ch
  if ch=='\\' and not escape:escape=True
  else:escape=False
 result.append(clean(current));return result

doc=Document();section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.65);section.left_margin=section.right_margin=Inches(.7)
for name,size in [('Normal',11),('Title',24),('Heading 1',17),('Heading 2',14),('Heading 3',12),('Heading 4',11)]:
 style=doc.styles[name];style.font.name='Arial';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
 style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
 style.paragraph_format.space_after=Pt(4);style.paragraph_format.line_spacing=1.15
for style in doc.styles:
 for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for name in ['Heading 1','Heading 2','Heading 3','Heading 4']:doc.styles[name].paragraph_format.space_before=Pt(13)
footer=section.footer.paragraphs[0];footer.alignment=2
r=footer.add_run('生鲜水产系统开发设计  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
lines=source.read_text(encoding='utf-8').splitlines();i=0;code=False
while i<len(lines):
 line=lines[i]
 if line.strip().startswith(('```','~~~')):code=not code;i+=1;continue
 if code:
  p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
  r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(9);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC');i+=1;continue
 if not line.strip():i+=1;continue
 if line.startswith('|') and i+1<len(lines) and re.match(r'^\|[ :|\-]+\|$',lines[i+1]):
  rows=[cells(line)];i+=2
  while i<len(lines) and lines[i].startswith('|'):rows.append(cells(lines[i]));i+=1
  count=max(map(len,rows));table=doc.add_table(rows=0,cols=count);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
  lengths=[max(min(len(row[n]) if n<len(row) else 0,35),8) for n in range(count) for row in [max(rows,key=lambda x:len(x[n]) if n<len(x) else 0)]]
  widths=[7.1*v/sum(lengths) for v in lengths]
  for n,w in enumerate(widths):table.columns[n].width=Inches(w)
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
  table._tbl.tblPr.append(borders)
  for ri,row in enumerate(rows):
   tr=table.add_row()
   if ri==0:header=OxmlElement('w:tblHeader');tr._tr.get_or_add_trPr().append(header)
   for n,cell in enumerate(tr.cells):
    cell.width=Inches(widths[n]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     v=OxmlElement('w:'+edge);v.set(qn('w:w'),'100');v.set(qn('w:type'),'dxa');margins.append(v)
    cell._tc.get_or_add_tcPr().append(margins);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
    r=p.add_run(row[n] if n<len(row) else '');r.font.size=Pt(10);r.bold=ri==0
    if ri==0:shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(shade)
  doc.add_paragraph();continue
 match=re.match(r'^(#{1,5})\s+(.+)',line)
 if match:
  depth=len(match[1]);text=re.sub(r'[^\w\s\u4e00-\u9fff]',' ',clean(match[2]))
  doc.add_paragraph(re.sub(r'\s+',' ',text).strip(),'Title' if depth==1 else 'Heading '+str(min(depth-1,4)));i+=1;continue
 if line.startswith('>'):doc.add_paragraph(clean(line.lstrip('> ')))
 elif line.lstrip().startswith('- '):doc.add_paragraph(clean(line.lstrip()[2:]),'List Bullet')
 elif re.match(r'^\d+\. ',line):doc.add_paragraph(clean(line))
 else:doc.add_paragraph(clean(line))
 i+=1
doc.core_properties.author='fishMall';doc.core_properties.title='生鲜水产线上经营系统开发设计文档';doc.core_properties.subject='一期业务设计与部署交付';doc.save(target)
print(target)
